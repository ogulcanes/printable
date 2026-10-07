from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ogulcan\Documents\GitHub\printable")
REFERENCE = ROOT / "output" / "Mega_Insulation_Figur_Fiyat_Teklifi.docx"
FINAL = ROOT / "output" / "Chakils_Borusan_Next_Fiyat_Teklifi.docx"

NAVY = "14213D"
ORANGE = "FF6542"
BLACK = "000000"
INK = "11141A"
GRAY = "6B7280"
PALE = "F1F3F6"
WHITE = "FFFFFF"
BORDER = "D9D9D9"
FONT = "Calibri"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, size: float, color: str = INK, bold: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
    alignment=None,
    before: float = 0,
    after: float = 0,
    line: float = 1.0,
):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size, color, bold)
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    return run


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_text(
    cell,
    text: str,
    *,
    size: float = 10.5,
    color: str = INK,
    bold: bool = False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size, color, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_bullet(paragraph, label: str, text: str) -> None:
    clear_paragraph(paragraph)
    bullet = paragraph.add_run("•  ")
    set_run_font(bullet, 10.5, ORANGE, True)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 10.5, INK, True)
    text_run = paragraph.add_run(" " + text)
    set_run_font(text_run, 10.5, INK, False)
    paragraph.paragraph_format.left_indent = Inches(0.20)
    paragraph.paragraph_format.first_line_indent = Inches(-0.20)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05


def set_grid_widths(table, widths: list[float]) -> None:
    table.autofit = False
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(Inches(width).twips))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def remove_paragraph_border(paragraph_or_style) -> None:
    element = getattr(paragraph_or_style, "_element", None)
    if element is None:
        element = paragraph_or_style._p
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def main() -> None:
    assert round(106.25 * 1000, 2) == 106250.00
    assert round(106250.00 * 0.20, 2) == 21250.00
    assert round(106250.00 * 1.20, 2) == 127500.00

    FINAL.parent.mkdir(parents=True, exist_ok=True)
    copy2(REFERENCE, FINAL)
    doc = Document(str(FINAL))

    # Styles required for semantic Word structure while retaining the source look.
    title_style = doc.styles["Title"]
    title_style.font.name = FONT
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = rgb(BLACK)
    title_style.paragraph_format.space_after = Pt(6)
    title_style.paragraph_format.line_spacing = 1.0
    remove_paragraph_border(title_style)

    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = FONT
    heading_style.font.size = Pt(14)
    heading_style.font.bold = True
    heading_style.font.color.rgb = rgb(BLACK)
    heading_style.paragraph_format.space_before = Pt(0)
    heading_style.paragraph_format.space_after = Pt(8)
    heading_style.paragraph_format.keep_with_next = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)

    # Preserve and label the source logo for accessibility.
    for shape in doc.inline_shapes:
        try:
            shape._inline.docPr.set("descr", "Printable logosu")
            shape._inline.docPr.set("title", "Printable")
        except Exception:
            pass

    # Header identity block.
    identity = doc.tables[0]
    right = identity.cell(0, 1)
    while len(right.paragraphs) < 3:
        right.add_paragraph()
    set_paragraph_text(
        right.paragraphs[0], "FİYAT TEKLİFİ", size=11, color=ORANGE, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT, after=8
    )
    set_paragraph_text(
        right.paragraphs[1], "Tarih 23.09.2026", size=10.5, color="D7DDEA",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT, after=6
    )
    set_paragraph_text(
        right.paragraphs[2], "Alıcı Chakil's Event Agency", size=10.5, color="D7DDEA",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT
    )

    # Main copy.
    paragraphs = doc.paragraphs
    paragraphs[1].style = doc.styles["Title"]
    remove_paragraph_border(paragraphs[1])
    set_paragraph_text(
        paragraphs[1], "Borusan Next Figürü Fiyat Teklifi", size=22, color=BLACK,
        bold=True, after=6
    )
    set_paragraph_text(
        paragraphs[2], "1.000 adet özel üretim", size=11, color=BLACK,
        after=10
    )
    set_paragraph_text(
        paragraphs[3],
        "Sayın Chakil's Event Agency Yetkilisi, talep edilen Borusan Next figürünün "
        "1.000 adet üretimi için hazırladığımız fiyat teklifini bilgilerinize sunarız.",
        size=10.5, color=INK, after=4, line=1.12
    )
    set_paragraph_text(paragraphs[4], "", size=4, after=0)
    paragraphs[5].style = doc.styles["Heading 1"]
    set_paragraph_text(paragraphs[5], "Fiyatlandırma", size=14, color=BLACK, bold=True, after=8)
    set_paragraph_text(paragraphs[6], "", size=2, after=0)
    paragraphs[7].style = doc.styles["Heading 1"]
    set_paragraph_text(paragraphs[7], "Ticari Koşullar", size=14, color=BLACK, bold=True, after=7)

    # Convert the existing pricing component to a four-column commercial table.
    table = doc.tables[1]
    if len(table.columns) == 3:
        table.add_column(Inches(1.72))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [2.75, 1.05, 1.40, 1.72]
    set_grid_widths(table, widths)
    set_table_borders(table)

    headers = ["Ürün ve kapsam", "Miktar", "Birim fiyat", "Tutar"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_text(
            cell, value, size=10, color=WHITE, bold=True,
            alignment=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        )

    product_values = [
        ("Borusan Next figürü üretimi", True, WD_ALIGN_PARAGRAPH.LEFT),
        ("1.000 adet", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("106,25 TL", False, WD_ALIGN_PARAGRAPH.RIGHT),
        ("106.250,00 TL", True, WD_ALIGN_PARAGRAPH.RIGHT),
    ]
    for index, (value, bold, alignment) in enumerate(product_values):
        cell = table.rows[1].cells[index]
        set_cell_shading(cell, WHITE)
        set_cell_text(cell, value, size=10.25, color=INK, bold=bold, alignment=alignment)

    summary_rows = [
        (2, "Ara toplam", "106.250,00 TL", PALE, INK),
        (3, "KDV (%20)", "21.250,00 TL", WHITE, INK),
        (4, "Genel toplam", "127.500,00 TL", NAVY, WHITE),
    ]
    for row_index, label, value, fill, text_color in summary_rows:
        row = table.rows[row_index]
        label_cell = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
        value_cell = row.cells[3]
        set_cell_shading(label_cell, fill)
        set_cell_shading(value_cell, fill)
        set_cell_text(
            label_cell, label, size=10.25, color=text_color,
            bold=row_index == 4, alignment=WD_ALIGN_PARAGRAPH.RIGHT
        )
        set_cell_text(
            value_cell, value, size=10.25, color=text_color,
            bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT
        )
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Inches(0.34)

    # Commercial conditions.
    set_bullet(
        paragraphs[8], "KDV", "Birim ve toplam fiyatlara KDV dahil değildir; hesaplamada %20 KDV esas alınmıştır."
    )
    set_bullet(
        paragraphs[9], "Ödeme", "Sipariş onayında teklif bedelinin %30'u ön ödeme olarak alınır. Kalan %70, üretim tamamlandığında ve sevkiyat öncesinde ödenir."
    )
    set_bullet(
        paragraphs[10], "Termin", "Ön ödemenin alınması ve üretim modelinin onaylanmasını takiben 2 aydır."
    )
    set_bullet(
        paragraphs[11], "Teklif geçerliliği", "Bu teklif 15 gün süreyle, 08.10.2026 tarihine kadar geçerlidir."
    )
    set_bullet(
        paragraphs[12], "Üretim onayı", "Ürün özellikleri, renk ve ölçüler üretim öncesi model veya numune onayıyla kesinleşir."
    )

    # Footer retained from source and expanded with verified company contact information.
    clear_paragraph(paragraphs[13])
    paragraphs[13].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[13].paragraph_format.space_before = Pt(10)
    paragraphs[13].paragraph_format.space_after = Pt(0)
    paragraphs[13].paragraph_format.line_spacing = 1.0
    footer_1 = paragraphs[13].add_run("Printable  |  Özel üretim ve 3D baskı çözümleri")
    set_run_font(footer_1, 8.5, GRAY, False)
    footer_1.add_break()
    footer_2 = paragraphs[13].add_run("0543 687 4208  |  info@printable.com.tr")
    set_run_font(footer_2, 8.5, GRAY, False)

    doc.core_properties.title = "Borusan Next Figürü Fiyat Teklifi"
    doc.core_properties.subject = "Chakil's Event Agency için 1.000 adet üretim teklifi"
    doc.core_properties.author = "Printable"
    doc.core_properties.keywords = "Borusan Next, figür üretimi, fiyat teklifi, Printable"
    doc.core_properties.modified = datetime(2026, 9, 23, 12, 0, 0)

    doc.save(str(FINAL))
    print(FINAL)


if __name__ == "__main__":
    main()
