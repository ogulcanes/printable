from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ogulcan\Documents\GitHub\printable")
LOGO = ROOT / "assets" / "printable-logo-transparent-400.png"
FINAL = ROOT / "output" / "Chakils_Borusan_Next_Fiyat_Teklifi_Revize.docx"

FONT = "Aptos"
NAVY = "182030"
ORANGE = "F85830"
BLACK = "000000"
INK = "1F2430"
GRAY = "667085"
LINE = "D9DEE7"
PALE = "F5F7FA"
PALE_ORANGE = "FFF3EE"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size: float, color: str = INK, bold: bool | None = None) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_paragraph_border(paragraph_or_style) -> None:
    element = getattr(paragraph_or_style, "_element", None) or paragraph_or_style._p
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def paragraph_text(
    paragraph,
    text: str,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 0,
    line: float = 1.0,
) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    run = paragraph.add_run(text)
    set_font(run, size, color, bold)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(
    cell,
    text: str,
    *,
    size: float = 10.25,
    color: str = INK,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_font(run, size, color, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color=LINE, size="6", include_inside=True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    edges = ["top", "left", "bottom", "right"]
    if include_inside:
        edges += ["insideH", "insideV"]
    for edge in edges:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_grid_widths(table, widths: list[float]) -> None:
    table.autofit = False
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(Inches(width).twips))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Inches(7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_grid_widths(table, [3.6, 3.6])
    remove_table_borders(table)
    left, right = table.rows[0].cells
    set_cell_text(left, "Printable  |  Özel üretim ve 3D baskı çözümleri", size=8.5, color=GRAY)
    set_cell_text(
        right, "0543 687 4208  |  info@printable.com.tr", size=8.5, color=GRAY,
        align=WD_ALIGN_PARAGRAPH.RIGHT
    )


def main() -> None:
    assert LOGO.exists()
    assert 106.25 * 1000 == 106250.0
    assert 106250.0 * 0.20 == 21250.0
    assert 106250.0 * 1.20 == 127500.0

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(0)

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = rgb(BLACK)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.line_spacing = 1.0
    remove_paragraph_border(title)

    heading = styles["Heading 1"]
    heading.font.name = FONT
    heading.font.size = Pt(13)
    heading.font.bold = True
    heading.font.color.rgb = rgb(BLACK)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True

    # Brand and offer metadata.
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_grid_widths(header, [4.35, 2.85])
    remove_table_borders(header)
    left, right = header.rows[0].cells
    set_cell_margins(left, top=0, start=0, bottom=40, end=80)
    set_cell_margins(right, top=0, start=80, bottom=40, end=0)
    logo_p = left.paragraphs[0]
    logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_p.paragraph_format.space_after = Pt(0)
    logo_run = logo_p.add_run()
    logo_run.add_picture(str(LOGO), width=Inches(2.35))
    for shape in doc.inline_shapes:
        shape._inline.docPr.set("descr", "Printable logosu")
        shape._inline.docPr.set("title", "Printable")

    right.text = ""
    offer_p = right.paragraphs[0]
    paragraph_text(
        offer_p, "FİYAT TEKLİFİ", size=10.5, color=ORANGE, bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT, after=6
    )
    meta_1 = right.add_paragraph()
    paragraph_text(meta_1, "23 Eylül 2026", size=9.5, color=INK, align=WD_ALIGN_PARAGRAPH.RIGHT, after=3)
    meta_2 = right.add_paragraph()
    paragraph_text(meta_2, "Chakil's Event Agency", size=9.5, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, after=3)
    meta_3 = right.add_paragraph()
    paragraph_text(meta_3, "Geçerlilik 15 gün", size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph(style="Title")
    remove_paragraph_border(p)
    paragraph_text(p, "Borusan Next Figürü Üretim Teklifi", size=22, color=BLACK, bold=True, after=5)
    p = doc.add_paragraph()
    paragraph_text(p, "1.000 adet özel üretim", size=11, color=BLACK, bold=True, after=10)
    p = doc.add_paragraph()
    paragraph_text(
        p,
        "Sayın Chakil's Event Agency Yetkilisi, talep edilen Borusan Next figürünün "
        "1.000 adet üretimi için hazırladığımız teklif aşağıda sunulmuştur.",
        size=10.5, color=INK, after=13, line=1.12
    )

    p = doc.add_paragraph(style="Heading 1")
    paragraph_text(p, "Fiyatlandırma", size=13, color=BLACK, bold=True, after=7)

    price = doc.add_table(rows=5, cols=4)
    price.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_grid_widths(price, [2.85, 1.05, 1.35, 1.95])
    set_table_borders(price)
    mark_header_row(price.rows[0])

    for i, value in enumerate(("Ürün ve kapsam", "Miktar", "Birim fiyat", "Tutar")):
        cell = price.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_text(
            cell, value, size=9.75, color=WHITE, bold=True,
            align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        )

    product = (
        ("Borusan Next figürü üretimi", True, WD_ALIGN_PARAGRAPH.LEFT),
        ("1.000 adet", False, WD_ALIGN_PARAGRAPH.CENTER),
        ("106,25 TL", False, WD_ALIGN_PARAGRAPH.RIGHT),
        ("106.250,00 TL", True, WD_ALIGN_PARAGRAPH.RIGHT),
    )
    for i, (value, bold, align) in enumerate(product):
        cell = price.rows[1].cells[i]
        set_cell_shading(cell, WHITE)
        set_cell_text(cell, value, size=10, color=INK, bold=bold, align=align)

    summary = (
        (2, "Ara toplam", "106.250,00 TL", PALE, INK),
        (3, "KDV (%20)", "21.250,00 TL", WHITE, INK),
        (4, "Genel toplam", "127.500,00 TL", PALE_ORANGE, NAVY),
    )
    for row_i, label, amount, fill, color in summary:
        row = price.rows[row_i]
        label_cell = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
        amount_cell = row.cells[3]
        set_cell_shading(label_cell, fill)
        set_cell_shading(amount_cell, fill)
        set_cell_text(label_cell, label, size=10, color=color, bold=row_i == 4, align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(amount_cell, amount, size=10, color=color, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for row in price.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Inches(0.35)

    p = doc.add_paragraph()
    paragraph_text(p, "Tüm tutarlar Türk Lirası cinsindendir.", size=8.5, color=GRAY, after=12)

    p = doc.add_paragraph(style="Heading 1")
    paragraph_text(p, "Ticari Koşullar", size=13, color=BLACK, bold=True, after=7)

    terms = doc.add_table(rows=5, cols=2)
    terms.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_grid_widths(terms, [1.55, 5.65])
    set_table_borders(terms, color=LINE, size="5")
    term_rows = (
        ("Ödeme", "Sipariş onayında teklif bedelinin %30'u ön ödeme olarak alınır. Kalan %70, üretim tamamlandığında ve sevkiyat öncesinde ödenir."),
        ("Termin", "Ön ödemenin alınması ve nihai üretim modelinin onaylanmasını takiben 2 aydır."),
        ("Geçerlilik", "Teklif 15 gün süreyle, 08.10.2026 tarihine kadar geçerlidir."),
        ("KDV", "Birim ve toplam fiyatlara KDV dahil değildir; hesaplamada %20 KDV esas alınmıştır."),
        ("Üretim onayı", "Ürün özellikleri, renk ve ölçüler üretim öncesi model veya numune onayıyla kesinleşir."),
    )
    for row_i, (label, value) in enumerate(term_rows):
        label_cell, value_cell = terms.rows[row_i].cells
        fill = PALE if row_i % 2 == 0 else WHITE
        set_cell_shading(label_cell, fill)
        set_cell_shading(value_cell, fill)
        set_cell_text(label_cell, label, size=9.75, color=ORANGE, bold=True)
        set_cell_text(value_cell, value, size=9.75, color=INK)
        terms.rows[row_i].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        terms.rows[row_i].height = Inches(0.38)

    add_footer(section)

    doc.core_properties.title = "Borusan Next Figürü Üretim Teklifi"
    doc.core_properties.subject = "Chakil's Event Agency için 1.000 adet üretim teklifi"
    doc.core_properties.author = "Printable"
    doc.core_properties.keywords = "Borusan Next, figür üretimi, fiyat teklifi, Printable"
    doc.core_properties.modified = datetime(2026, 9, 23, 12, 0, 0)

    FINAL.parent.mkdir(parents=True, exist_ok=True)
    doc.save(FINAL)
    print(FINAL)


if __name__ == "__main__":
    main()
