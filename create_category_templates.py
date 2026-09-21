from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import json, os, shutil
from urllib.request import urlopen
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / 'outputs' / 'trendyol-kategori-sablonlari'
ZIP_PATH = ROOT / 'outputs' / 'trendyol-kategori-sablonlari.zip'
PRODUCTS = json.loads((Path(os.environ['TEMP']) / 'printable-products.json').read_text(encoding='utf-8-sig'))

GROUPS = [
 (833,'Figur',[50,29,28,26,21,19,17,12,10,6],{}),(2840,'Anahtarlik',[54,38,36,33,24,22,18,13,9,5],{'Renk':'Renk','Web Color':'Web Color','Kutu Durumu':'Kutu yok'}),
 (1015,'Diger_Oyuncaklar',[53,52,51,49,48,47,46,45,44,43,42,41,40,39,35,16,7],{'Yaş':'10+ Yaş','Renk':'Renk','Web Color':'Web Color'}),(3618,'Duduk',[15,14,11],{'Renk':'Renk','Web Color':'Web Color'}),
 (4937,'Tuvalet_Kagitligi',[37],{'Renk':'Renk','Web Color':'Web Color','Materyal':'Plastik'}),(5468,'Diger_Yazici_Sarf',[34],{'Renk':'Renk','Web Color':'Web Color'}),(1511,'Kalemlik',[30],{'Materyal':'Plastik'}),
 (4973,'Makyaj_Aplikatorleri',[27],{'Renk':'Renk','Web Color':'Web Color'}),(5206,'Sise_Acacagi',[31],{'Renk':'Renk','Web Color':'Web Color'}),(1881,'Vazo',[23],{'Materyal':'Plastik','Yükseklik':'Belirtilmemiş','Parça Sayısı':'1 Parça','Renk':'Renk','Web Color':'Web Color'}),(2710,'Tepsi',[20],{'Parça Sayısı':'1 Parça','Materyal':'Plastik','Web Color':'Web Color','Boyut/Ebat':'Belirtilmemiş','Renk':'Renk'})]

BASE_HEADERS=['Barkod','Model Kodu','Marka','Kategori','Para Birimi','Ürün Adı','Ürün Açıklaması','Piyasa Satış Fiyatı (KDV Dahil)',"Trendyol'da Satılacak Fiyat (KDV Dahil)",'Ürün Stok Adedi','Stok Kodu','KDV Oranı','ÖTV Oranı','Desi','Parti/Lot/SKT Bilgisi','Görsel 1','Görsel 2','Görsel 3','Görsel 4','Görsel 5','Görsel 6','Görsel 7','Görsel 8','Sevkiyat Süresi','Sevkiyat Tipi','Birincil İthalatçı Mail Adresi','Diğer Özellikler','Üretici Adı','Birincil İthalatçı Adı','İkincil İthalatçı Mail Adresi','Birincil İthalatçı Adres Bilgisi','Üçüncül İthalatçı Adres Bilgisi','Üretici Mail Adresi','Üretici Adres Bilgisi','Menşei','İkincil İthalatçı Adı','Üçüncül İthalatçı Mail Adresi','İkincil İthalatçı Adres Bilgisi','Üçüncül İthalatçı Adı','Boyut']
PRODUCER_NAME='Printable 3D Baskı Atölyesi'; PRODUCER_EMAIL='info@printable.com.tr'; PRODUCER_ADDRESS='Emniyetevleri Mahallesi, Bülbüldere Sokak No: 12, Kağıthane/İstanbul'

def ean(pid):
    base=f'869{pid:09d}'; return base+str((10-sum(int(d)*(1 if i%2==0 else 3) for i,d in enumerate(base))%10)%10)
def price(p):
    scales=p.get('scales') or []
    return min((Decimal(str(x['price'])) for x in scales),default=Decimal(str(p.get('sale_price') or p['price'])))
def color(p):
    colors=[x.get('name') for x in (p.get('colors') or []) if x.get('name')]
    return colors[0] if len(colors)==1 else 'Çok Renkli'
def images(p):
    result=[]
    for value in [p.get('image_path')]+[x.get('image_path') for x in (p.get('images') or [])]:
        if value and value not in result: result.append(value if value.startswith('http') else 'https://www.printable.com.tr'+value)
    return result[:8]
def category_attributes(category_id):
    with urlopen(f'https://apigw.trendyol.com/integration/product/categories/{category_id}/attributes',timeout=30) as response:
        data=json.load(response)
    return [item['attribute']['name'] for item in data.get('categoryAttributes',[]) if item['attribute']['name'] not in BASE_HEADERS]

if OUT_DIR.exists(): shutil.rmtree(OUT_DIR)
OUT_DIR.mkdir(parents=True)
header_fill=PatternFill('solid',fgColor='FCE4D6'); header_font=Font(bold=True,color='FF0000'); header_align=Alignment(horizontal='center',vertical='center')
created=[]
for category_id,slug,ids,values in GROUPS:
    attrs=category_attributes(category_id)
    headers=BASE_HEADERS+attrs
    wb=Workbook(); ws=wb.active; ws.title=f'{slug}_{category_id}'[:31]; ws.append(headers); ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).coordinate}'
    for cell in ws[1]: cell.fill=header_fill; cell.font=header_font; cell.alignment=header_align
    for product in [p for p in PRODUCTS if p['id'] in ids]:
        final=(price(product)*Decimal('1.22')).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP); model=product.get('sku') or f"PRINTABLE-{product['id']}"
        row=[ean(product['id']),model,'',category_id,'TRY',product['name'],product.get('description') or '',float(final),float(final),int(product.get('stock') or 0),model,20,0,1,'']+images(product)+['']*(8-len(images(product)))+[1,'Hızlı Teslimat','','3D baskı PLA plastik ürün',PRODUCER_NAME,'','','','',PRODUCER_EMAIL,PRODUCER_ADDRESS,'TR','','','','','']
        for attr in attrs:
            value=values.get(attr,'')
            row.append(color(product) if value in ('Renk','Web Color') else value)
        ws.append(row)
    ws.column_dimensions['A'].width=16; ws.column_dimensions['B'].width=18; ws.column_dimensions['F'].width=38; ws.column_dimensions['G'].width=70
    for row in ws.iter_rows(min_row=2):
        row[0].number_format='@'; row[7].number_format='0.00'; row[8].number_format='0.00'
    path=OUT_DIR / f'{category_id}_{slug}.xlsx'; wb.save(path)
    check=load_workbook(path,read_only=True); assert check.active.max_row-1==len(ids); check.close(); created.append(path)

if ZIP_PATH.exists(): ZIP_PATH.unlink()
shutil.make_archive(str(ZIP_PATH.with_suffix('')),'zip',OUT_DIR)
print(f'files={len(created)} products={sum(len(group[2]) for group in GROUPS)} zip={ZIP_PATH}')
