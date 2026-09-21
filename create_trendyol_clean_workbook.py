from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import json, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'outputs' / 'trendyol-urun-yukleme-2026-09-20-tek-sekme-sade.xlsx'
PRODUCTS = json.loads((Path(os.environ['TEMP']) / 'printable-products.json').read_text(encoding='utf-8-sig'))

GROUPS = [
 (833,'Figur_833',[50,29,28,26,21,19,17,12,10,6],{}),(2840,'Anahtarlik_2840',[54,38,36,33,24,22,18,13,9,5],{'Renk':'Renk','Web Color':'Web Color','Kutu Durumu':'Kutu yok'}),
 (1015,'Diger_Oyuncaklar_1015',[53,52,51,49,48,47,46,45,44,43,42,41,40,39,35,16,7],{'Yaş':'10+ Yaş','Renk':'Renk','Web Color':'Web Color'}),(3618,'Duduk_3618',[15,14,11],{'Renk':'Renk','Web Color':'Web Color'}),
 (4937,'Tuvalet_Kagitligi_4937',[37],{'Renk':'Renk','Web Color':'Web Color','Materyal':'Plastik'}),(5468,'Diger_Yazici_Sarf_5468',[34],{'Renk':'Renk','Web Color':'Web Color'}),(1511,'Kalemlik_1511',[30],{'Materyal':'Plastik'}),
 (4973,'Makyaj_Aplikatorleri_4973',[27],{'Renk':'Renk','Web Color':'Web Color'}),(5206,'Sise_Acacagi_5206',[31],{'Renk':'Renk','Web Color':'Web Color'}),(1881,'Vazo_1881',[23],{'Materyal':'Plastik','Yükseklik':'Belirtilmemiş','Parça Sayısı':'1 Parça','Renk':'Renk','Web Color':'Web Color'}),(2710,'Tepsi_2710',[20],{'Parça Sayısı':'1 Parça','Materyal':'Plastik','Web Color':'Web Color','Boyut/Ebat':'Belirtilmemiş','Renk':'Renk'})]

HEADERS=['Barkod','Model Kodu','Marka','Kategori','Para Birimi','Ürün Adı','Ürün Açıklaması','Piyasa Satış Fiyatı (KDV Dahil)',"Trendyol'da Satılacak Fiyat (KDV Dahil)",'Ürün Stok Adedi','Stok Kodu','KDV Oranı','ÖTV Oranı','Desi','Parti/Lot/SKT Bilgisi','Görsel 1','Görsel 2','Görsel 3','Görsel 4','Görsel 5','Görsel 6','Görsel 7','Görsel 8','Sevkiyat Süresi','Sevkiyat Tipi','Birincil İthalatçı Mail Adresi','Diğer Özellikler','Üretici Adı','Birincil İthalatçı Adı','İkincil İthalatçı Mail Adresi','Birincil İthalatçı Adres Bilgisi','Üçüncül İthalatçı Adres Bilgisi','Üretici Mail Adresi','Üretici Adres Bilgisi','Menşei','İkincil İthalatçı Adı','Üçüncül İthalatçı Mail Adresi','İkincil İthalatçı Adres Bilgisi','Üçüncül İthalatçı Adı','Boyut']
PRODUCER_NAME='Printable 3D Baskı Atölyesi'
PRODUCER_EMAIL='info@printable.com.tr'
PRODUCER_ADDRESS='Emniyetevleri Mahallesi, Bülbüldere Sokak No: 12, Kağıthane/İstanbul'

def ean(pid):
    base=f'869{pid:09d}'; return base+str((10-sum(int(d)*(1 if i%2==0 else 3) for i,d in enumerate(base))%10)%10)
def price(p):
    scales=p.get('scales') or []
    return min((Decimal(str(x['price'])) for x in scales),default=Decimal(str(p.get('sale_price') or p['price'])))
def color(p):
    vals=[x.get('name') for x in (p.get('colors') or []) if x.get('name')]
    return vals[0] if len(vals)==1 else 'Çok Renkli'
def images(p):
    paths=[p.get('image_path')]+[x.get('image_path') for x in (p.get('images') or [])]
    result=[]
    for x in paths:
        if x and x not in result: result.append(x if x.startswith('http') else 'https://www.printable.com.tr'+x)
    return result[:8]

wb=Workbook(); ws=wb.active; ws.title='Ürünler'
header_fill=PatternFill('solid',fgColor='FCE4D6'); header_font=Font(bold=True,color='FF0000'); header_align=Alignment(horizontal='center',vertical='center')
headers=HEADERS; ws.append(headers); ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).coordinate}'
for cell in ws[1]: cell.fill=header_fill; cell.font=header_font; cell.alignment=header_align
for category_id,name,ids,attrs in GROUPS:
    for p in [x for x in PRODUCTS if x['id'] in ids]:
        final=(price(p)*Decimal('1.22')).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP); model=p.get('sku') or f"PRINTABLE-{p['id']}"
        row=[ean(p['id']),model,'',category_id,'TRY',p['name'],p.get('description') or '',float(final),float(final),int(p.get('stock') or 0),model,20,0,1,'']+images(p)+['']*(8-len(images(p)))+[1,'Hızlı Teslimat','','3D baskı PLA plastik ürün',PRODUCER_NAME,'','','','',PRODUCER_EMAIL,PRODUCER_ADDRESS,'TR','','','','','']
        ws.append(row)
ws.column_dimensions['A'].width=16; ws.column_dimensions['B'].width=18; ws.column_dimensions['F'].width=38; ws.column_dimensions['G'].width=70
for col in range(8,10):
    for cell in list(ws.columns)[col-1][1:]: cell.number_format='0.00'
for cell in list(ws.columns)[0][1:]: cell.number_format='@'
OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
check=load_workbook(OUT,read_only=True); assert check.sheetnames==['Ürünler']; assert check.active.max_row-1==47; check.close(); print(OUT)
