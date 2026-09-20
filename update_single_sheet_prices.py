from decimal import Decimal
from pathlib import Path
import json, os, shutil
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'outputs' / 'trendyol-urun-yukleme-2026-09-20-tek-sekme-sade.xlsx'
OUTPUT = ROOT / 'outputs' / 'trendyol-urun-yukleme-2026-09-20-tek-sekme-fiyat-guncel.xlsx'
PRODUCTS = json.loads((Path(os.environ['TEMP']) / 'printable-products.json').read_text(encoding='utf-8-sig'))

def effective_price(product):
    scales = product.get('scales') or []
    if scales:
        return min(Decimal(str(item['price'])) for item in scales)
    return Decimal(str(product.get('sale_price') or product['price']))

def rounded_price(value):
    # Bir sonraki ...9 TL seviyesine çıkarılır; komisyon sonrası marj korunur.
    whole = int(value.to_integral_value(rounding='ROUND_CEILING'))
    return whole if whole % 10 == 9 else whole + (9 - whole % 10)

by_name = {product['name']: product for product in PRODUCTS}
shutil.copy2(SOURCE, OUTPUT)
workbook = load_workbook(OUTPUT)
sheet = workbook['Ürünler']
updated = 0
for row in range(2, sheet.max_row + 1):
    name = sheet.cell(row, 6).value
    product = by_name.get(name)
    if not product:
        raise ValueError(f'Canlı sitede bulunamayan ürün: {name}')
    price = rounded_price(effective_price(product) * Decimal('1.22'))
    sheet.cell(row, 8).value = price
    sheet.cell(row, 9).value = price
    sheet.cell(row, 8).number_format = '0'
    sheet.cell(row, 9).number_format = '0'
    updated += 1
workbook.save(OUTPUT)

check = load_workbook(OUTPUT, read_only=True)
assert check.sheetnames == ['Ürünler'] and check.active.max_row - 1 == 47 and updated == 47
assert all(isinstance(check.active.cell(row, 8).value, int) and check.active.cell(row, 8).value % 10 == 9 for row in range(2, 49))
check.close()
print(OUTPUT)
