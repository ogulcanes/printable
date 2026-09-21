from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import json
import os
import openpyxl

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / 'urun-olusturma-1416362.xlsx'
OUTPUT_DIR = ROOT / 'outputs'
OUTPUT = OUTPUT_DIR / 'trendyol-urun-yukleme-2026-09-20-duzeltilmis.xlsx'
PRODUCTS = json.loads((Path(os.environ['TEMP']) / 'printable-products.json').read_text(encoding='utf-8-sig'))

PRODUCT_GROUPS = [
    (833, 'Figür(833)', [50, 29, 28, 26, 21, 19, 17, 12, 10, 6], {}),
    (2840, 'Anahtarlık(2840)', [54, 38, 36, 33, 24, 22, 18, 13, 9, 5], {'Renk': 'Renk', 'Web Color': 'Web Color', 'Kutu Durumu': 'Kutu yok'}),
    (1015, 'Diger Oyuncaklar(1015)', [53, 52, 51, 49, 48, 47, 46, 45, 44, 43, 42, 41, 40, 39, 35, 16, 7], {'Yaş': '10+ Yaş', 'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (3618, 'Duduk(3618)', [15, 14, 11], {'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (4937, 'Tuvalet Kagitligi(4937)', [37], {'Renk': 'Renk', 'Web Color': 'Web Color', 'Materyal': 'Plastik'}),
    (5468, 'Diger Yazici Sarf(5468)', [34], {'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (1511, 'Kalemlik(1511)', [30], {'Materyal': 'Plastik'}),
    (4973, 'Makyaj Aplikatorleri(4973)', [27], {'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (5206, 'Sise Acacagi(5206)', [31], {'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (1881, 'Vazo(1881)', [23], {'Materyal': 'Plastik', 'Yükseklik': 'Belirtilmemiş', 'Parça Sayısı': '1 Parça', 'Renk': 'Renk', 'Web Color': 'Web Color'}),
    (2710, 'Tepsi(2710)', [20], {'Parça Sayısı': '1 Parça', 'Materyal': 'Plastik', 'Web Color': 'Web Color', 'Boyut/Ebat': 'Belirtilmemiş', 'Renk': 'Renk'}),
]

def ean13(product_id: int) -> str:
    base = f'869{product_id:09d}'
    total = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(base))
    return base + str((10 - total % 10) % 10)

def effective_price(product: dict) -> Decimal:
    scales = product.get('scales') or []
    if scales:
        return min(Decimal(str(scale['price'])) for scale in scales)
    return Decimal(str(product.get('sale_price') or product['price']))

def product_color(product: dict) -> str:
    colors = [color.get('name') for color in product.get('colors') or [] if color.get('name')]
    return colors[0] if len(colors) == 1 else 'Çok Renkli'

def product_images(product: dict) -> list[str]:
    paths = [product.get('image_path')] + [image.get('image_path') for image in product.get('images') or []]
    result = []
    for path in paths:
        if not path or path in result:
            continue
        result.append(path if path.startswith('http') else f'https://www.printable.com.tr{path}')
    return result[:8]

def clear_data_rows(ws):
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for cell in row:
            cell.value = None

def set_header(ws, column, value):
    cell = ws.cell(row=1, column=column, value=value)
    cell.font = copy(ws['A1'].font)
    cell.fill = copy(ws['A1'].fill)
    cell.alignment = copy(ws['A1'].alignment)
    cell.border = copy(ws['A1'].border)

def fill_sheet(ws, category_id, product_ids, attributes):
    clear_data_rows(ws)
    ws['D2'] = category_id
    ws['E2'] = 'TRY'
    selected = [product for product in PRODUCTS if product['id'] in product_ids]
    attr_columns = {}
    if category_id not in (833, 2840):
        for offset, attr_name in enumerate(attributes, start=41):
            set_header(ws, offset, attr_name)
            attr_columns[attr_name] = offset

    for row, product in enumerate(selected, start=3):
        site_price = effective_price(product)
        trendyol_price = (site_price * Decimal('1.22')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        model = product.get('sku') or f"PRINTABLE-{product['id']}"
        data = {
            1: ean13(product['id']), 2: model, 4: category_id, 5: 'TRY', 6: product['name'],
            7: product.get('description') or '', 8: float(trendyol_price), 9: float(trendyol_price),
            10: int(product.get('stock') or 0), 11: model, 12: 20, 13: 0, 14: 1,
            24: 3, 25: 'Hızlı Teslimat', 35: 'TR'
        }
        for index, image in enumerate(product_images(product), start=16):
            data[index] = image
        if category_id == 2840:
            data.update({26: product_color(product), 33: 'TR', 41: product_color(product), 43: 'Kutu yok'})
        for attr_name, attr_value in attributes.items():
            if category_id == 2840:
                continue
            value = product_color(product) if attr_value in ('Renk', 'Web Color') else attr_value
            data[attr_columns[attr_name]] = value
        for column, value in data.items():
            cell = ws.cell(row=row, column=column, value=value)
            if column == 1:
                cell.number_format = '@'
            elif column in (8, 9):
                cell.number_format = '0.00'

OUTPUT_DIR.mkdir(exist_ok=True)
workbook = openpyxl.load_workbook(TEMPLATE)
figure_sheet = workbook['Figür(833)']
keychain_sheet = workbook['Anahtarlık(2840)']

for category_id, sheet_name, product_ids, attributes in PRODUCT_GROUPS:
    if category_id == 833:
        sheet = figure_sheet
    elif category_id == 2840:
        sheet = keychain_sheet
    else:
        sheet = workbook.copy_worksheet(figure_sheet)
        sheet.title = sheet_name
    fill_sheet(sheet, category_id, product_ids, attributes)

workbook.save(OUTPUT)

# Ensures the file can be read again by a standards-compliant XLSX reader.
verification = openpyxl.load_workbook(OUTPUT, read_only=True, data_only=False)
assert len(verification.sheetnames) == 16
assert sum(1 for ws in verification.worksheets if ws.max_row >= 3) >= 11
verification.close()
print(OUTPUT)
